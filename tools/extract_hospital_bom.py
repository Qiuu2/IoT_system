# -*- coding: utf-8 -*-
"""从全专业清单抽取子系统构成与接入规模。

    python3 tools/extract_hospital_bom.py 全专业清单.xlsx docs/
"""
import csv
import re
import sys
from collections import OrderedDict
from pathlib import Path

import openpyxl

# 只有这些单位才是「一件设备」；米/条/项等按量计费，不进接入规模。
COUNT_UNITS = {'台', '套', '个', '只', '部', '块', '张', '面'}


def parse(xlsx):
    ws = openpyxl.load_workbook(xlsx, data_only=True).worksheets[0]
    rows = []
    for r in range(5, ws.max_row + 1):
        def g(c):
            v = ws.cell(r, c).value
            return '' if v is None else str(v).strip()
        name = g(2)
        if not name:
            continue
        try:
            qty = float(g(5))
        except ValueError:
            qty = 0.0
        rows.append({'row': r, 'seq': g(1), 'name': name, 'spec': g(3),
                     'unit': g(4), 'qty': qty, 'subsystem': g(8) or '（未标注）'})
    return rows


def write_csv(rows, out):
    cols = ['row', 'subsystem', 'seq', 'name', 'unit', 'qty']
    with open(out, 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.DictWriter(f, fieldnames=cols, extrasaction='ignore')
        w.writeheader()
        w.writerows(rows)


def write_summary(rows, out):
    sub = OrderedDict()
    for x in rows:
        sub.setdefault(x['subsystem'], []).append(x)

    def countable(items):
        return int(sum(i['qty'] for i in items if i['unit'] in COUNT_UNITS))

    with open(out, 'w', encoding='utf-8') as f:
        f.write('# 全专业清单：子系统构成\n\n')
        f.write('> 由 `tools/extract_hospital_bom.py` 自清单抽取。\n')
        f.write('> 「计件设备」只统计单位为 台/套/个/只/部/块/张/面 的行，'
                '已排除线缆等按米、按项计量的条目。\n\n')
        f.write('| 子系统 | 条目数 | 计件设备 |\n|---|---:|---:|\n')
        for k, v in sorted(sub.items(), key=lambda kv: -countable(kv[1])):
            f.write('| %s | %d | %s |\n' % (k, len(v), countable(v) or '—'))
        f.write('\n合计 %d 个子系统、%d 行、计件设备 %d 件。\n'
                % (len(sub), len(rows), countable(rows)))


def write_platform_spec(rows, out):
    """把『智能化集成平台』条目单独导出——那是本平台自身的招标规格。"""
    sel = [x for x in rows if x['subsystem'] == '智能化集成平台']
    with open(out, 'w', encoding='utf-8') as f:
        f.write('# 本平台的招标规格（清单原文）\n\n')
        f.write('> 自清单「智能化集成平台」子系统抽取，**未经改写**，'
                '是需求与验收的直接依据。\n\n')
        for x in sel:
            f.write('\n## %s · %s %s　`清单第 %d 行`\n\n'
                    % (x['seq'], x['name'], x['qty'] and
                       '（%g %s）' % (x['qty'], x['unit']) or '', x['row']))
            for line in x['spec'].splitlines():
                if line.strip():
                    f.write('%s\n' % line.strip())
            f.write('\n')
    return len(sel)


def main():
    xlsx, outdir = sys.argv[1], Path(sys.argv[2])
    outdir.mkdir(parents=True, exist_ok=True)
    rows = parse(xlsx)
    write_csv(rows, outdir / '全专业清单.csv')
    write_summary(rows, outdir / '子系统构成.md')
    n = write_platform_spec(rows, outdir / '平台招标规格-清单原文.md')
    print('设备行 %d · 子系统 %d · 平台自身条目 %d' % (
        rows and len(rows) or 0, len({x['subsystem'] for x in rows}), n))


if __name__ == '__main__':
    main()
